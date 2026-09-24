#!/usr/bin/env python3
"""Builds a resume or cover letter .docx from a JSON spec."""
import argparse
import json
import subprocess
import sys


def _ensure_python_docx():
    try:
        import docx  # noqa: F401
        return
    except ImportError:
        pass
    subprocess.check_call([sys.executable, "-m", "pip", "install", "--user", "python-docx"])


def _set_base_style(document):
    style = document.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)


def build_resume(data, output_path):
    document = Document()
    _set_base_style(document)

    name = document.add_paragraph()
    run = name.add_run(data["name"])
    run.bold = True
    run.font.size = Pt(20)

    if data.get("contact"):
        contact = document.add_paragraph(data["contact"])
        contact.runs[0].font.size = Pt(10)
        contact.runs[0].font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    if data.get("summary"):
        document.add_paragraph()
        document.add_paragraph(data["summary"])

    for section in data.get("sections", []):
        document.add_paragraph()
        heading = document.add_heading(section["heading"], level=2)
        for r in heading.runs:
            r.font.color.rgb = RGBColor(0x22, 0x22, 0x22)

        for entry in section.get("entries", []):
            entry_p = document.add_paragraph()
            title_run = entry_p.add_run(entry["title"])
            title_run.bold = True
            if entry.get("subtitle"):
                sub_run = entry_p.add_run("  |  " + entry["subtitle"])
                sub_run.italic = True
                sub_run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
            for bullet in entry.get("bullets", []):
                document.add_paragraph(bullet, style="List Bullet")

        for bullet in section.get("bullets", []):
            document.add_paragraph(bullet, style="List Bullet")

    document.save(output_path)


def build_letter(data, output_path):
    document = Document()
    _set_base_style(document)

    if data.get("sender_name"):
        document.add_paragraph(data["sender_name"])
    if data.get("sender_contact"):
        document.add_paragraph(data["sender_contact"])
    document.add_paragraph()
    if data.get("date"):
        document.add_paragraph(data["date"])
    document.add_paragraph()

    if data.get("recipient"):
        for line in data["recipient"].split("\n"):
            document.add_paragraph(line)
        document.add_paragraph()

    if data.get("salutation"):
        document.add_paragraph(data["salutation"])
        document.add_paragraph()

    for paragraph in data.get("paragraphs", []):
        document.add_paragraph(paragraph)
        document.add_paragraph()

    if data.get("closing"):
        document.add_paragraph(data["closing"])
    if data.get("signature_name"):
        document.add_paragraph()
        document.add_paragraph(data["signature_name"])

    document.save(output_path)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("kind", choices=["resume", "letter"])
    parser.add_argument("--data", required=True, help="Path to the JSON spec")
    parser.add_argument("--output", required=True, help="Path to write the .docx to")
    args = parser.parse_args()

    _ensure_python_docx()
    global Document, Pt, RGBColor
    from docx import Document
    from docx.shared import Pt, RGBColor

    with open(args.data, "r", encoding="utf-8") as f:
        data = json.load(f)

    if args.kind == "resume":
        build_resume(data, args.output)
    else:
        build_letter(data, args.output)

    print(f"Wrote {args.output}")


if __name__ == "__main__":
    main()
